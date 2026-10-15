"""VesselVio's side (Bumgarner & Nelson 2022, GPL-3): its own output files.

VesselVio is GPL-3 and HaemoLynx is Apache-2.0, so nothing here imports or
runs it: VesselVio is run as the program it is -- its app, or its
``VVTerminal.py`` with "save segment results" on -- and this reads the two
files it writes, the dataset workbook ``VesselVio 3D Dataset Analysis
Results.xlsx`` and ``Segment Results/<name>.csv``.
"""
from __future__ import annotations

import csv
import math
from pathlib import Path

import numpy as np

from .measures import Measurements

#: The workbook's columns this reads (row 2 is its header).
NETWORK_LENGTH = "Network Length"
BRANCHPOINTS = "Branchpoints"
ENDPOINTS = "Endpoints"
SEGMENT_COUNT = "Number of Segments"

DEFINITIONS = {
    "skeleton": "VesselVio's own Lee 1994 thinning; its graph's junction cliques filtered; end segments under `prune` um pruned and isolated segments under `filter` um removed (its terminal defaults: 5 and 10 um)",
    "segment": "a run of skeleton vertices between branch points or end points",
    "length": "arc length of a B-spline fitted through the segment's vertices (centreline smoothing on)",
    "tortuosity": "length over the chord between the smoothed segment's ends; a loop is 0",
    "diameter": "twice the segment's mean radius, each vertex's radius the mean distance to its four nearest background voxels' edges (not centres)",
    "branch point": "vertex of degree > 2 after clique filtering",
    "end point": "vertex of degree 1",
}


def read_segments(path: Path) -> list[dict[str, float]]:
    """The rows of a VesselVio ``Segment Results`` CSV (its first line names the file)."""
    with Path(path).open(newline="", encoding="utf-8") as f:
        lines = [line for line in f if line.strip()]
    rows = csv.DictReader(lines[1:])
    return [{k: float(v) for k, v in row.items() if k} for row in rows]


def read_network(path: Path, name: str | None = None) -> dict[str, float]:
    """The dataset row of a VesselVio results workbook (the last one, or the
    one whose ``File Name`` is *name*)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    try:
        rows = [[cell.value for cell in row] for row in workbook["Main Results"].iter_rows()]
    finally:
        workbook.close()
    header = rows[1]
    data = [dict(zip(header, row)) for row in rows[2:] if row and row[0] not in (None, "")]
    if name is not None:
        data = [row for row in data if row.get("File Name") == name]
    if not data:
        raise ValueError(f"No dataset row{f' for {name!r}' if name else ''} in {path}")
    return data[-1]


def measurements(segments: list[dict[str, float]], network: dict[str, float], *, tool: str = "VesselVio") -> Measurements:
    lengths = np.asarray([s["Length"] for s in segments], dtype=float)
    tortuosity = np.asarray([s["Tortuosity"] for s in segments], dtype=float)
    loops = int(np.count_nonzero(tortuosity == 0))
    total = float(network[NETWORK_LENGTH])
    if not math.isclose(total, float(lengths.sum()), rel_tol=1e-6):
        raise ValueError(f"segments sum to {lengths.sum()} um, the workbook says {total} um")
    return Measurements(
        tool=tool,
        total_length_um=total,
        segments=int(network[SEGMENT_COUNT]),
        branch_points=int(network[BRANCHPOINTS]),
        end_points=int(network[ENDPOINTS]),
        segment_lengths_um=lengths,
        tortuosity=np.where(tortuosity > 0, tortuosity, np.nan),
        diameters_um=2.0 * np.asarray([s["Mean Radius"] for s in segments], dtype=float),
        loops=loops,
        definitions=dict(DEFINITIONS),
    )


def load(results_dir: Path, name: str) -> Measurements:
    """VesselVio's measurements of the dataset *name* (its file stem) from a
    results folder."""
    results_dir = Path(results_dir)
    return measurements(
        read_segments(results_dir / "Segment Results" / f"{name}.csv"),
        read_network(results_dir / "VesselVio 3D Dataset Analysis Results.xlsx", name),
    )
